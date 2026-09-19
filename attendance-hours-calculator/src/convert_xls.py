# -*- coding: utf-8 -*-
"""Convert attendance .xls (multi-employee sheets) to matrix .xlsx.

Input: original machine-export .xls with sheets like Summary, Logs, 5.21.29, ...
Output: matrix .xlsx with one row per employee and date columns 07/01..07/NN.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import openpyxl
import pandas as pd
import xlrd


def parse_period(period_text: str) -> tuple:
    """Extract start/end day from '2026/07/01 ~ 07/26'. Returns (start_day, end_day)."""
    m = re.search(r'(\d{4})/(\d{2})/(\d{2})\s*~\s*(\d{2})/(\d{2})', period_text)
    if m:
        start_day = int(m.group(3))
        end_day = int(m.group(5))
        return start_day, end_day
    # fallback: try to find two day numbers
    nums = [int(x) for x in re.findall(r'/(\d{2})', period_text)]
    if len(nums) >= 2:
        return nums[0], nums[-1]
    raise ValueError(f'Cannot parse period: {period_text}')


def extract_employees_from_logs(wb: xlrd.book.Book) -> tuple:
    """Return (days, employees) where employees is {emp_no: {name, dept, daily}}."""
    logs = wb.sheet_by_name('Logs')

    # Day numbers are in row 3 (0-based)
    day_numbers = []
    for col in range(logs.ncols):
        v = logs.cell_value(3, col)
        if v != '' and isinstance(v, (int, float)):
            day_numbers.append(int(v))

    employees: Dict[str, dict] = {}
    row_idx = 4
    while row_idx < logs.nrows:
        header_row = logs.row_values(row_idx)
        if header_row and str(header_row[0]).strip() == 'No :':
            emp_no = str(header_row[2]).strip()
            emp_name = str(header_row[10]).strip()
            emp_dept = str(header_row[20]).strip()

            # Next row contains the punch data
            data_row = logs.row_values(row_idx + 1) if row_idx + 1 < logs.nrows else []
            daily: Dict[int, List[str]] = {}
            for col_idx, day in enumerate(day_numbers):
                if col_idx >= len(data_row):
                    continue
                val = data_row[col_idx]
                if val != '':
                    times = [t.strip() for t in str(val).strip().split('\n') if t.strip()]
                    if times:
                        daily[day] = times

            employees[emp_no] = {
                'name': emp_name,
                'dept': emp_dept,
                'daily': daily,
            }
            row_idx += 2
        else:
            row_idx += 1

    return day_numbers, employees


def extract_summary(wb: xlrd.book.Book) -> Dict[str, dict]:
    """Return {emp_no: summary_info} from Summary sheet."""
    summary = wb.sheet_by_name('Summary')
    result = {}
    for row_idx in range(4, summary.nrows):
        row = summary.row_values(row_idx)
        if not row or row[0] == '':
            continue
        emp_no = str(row[0]).strip()
        result[emp_no] = {
            'name': str(row[1]).strip() if len(row) > 1 else '',
            'dept': str(row[2]).strip() if len(row) > 2 else '',
            'required_hours': row[3] if len(row) > 3 else '',
            'actual_hours': row[4] if len(row) > 4 else '',
            'late_times': row[5] if len(row) > 5 else '',
            'late_min': row[6] if len(row) > 6 else '',
            'early_times': row[7] if len(row) > 7 else '',
            'early_min': row[8] if len(row) > 8 else '',
            'attend_req_act': row[11] if len(row) > 11 else '',
            'ab': row[13] if len(row) > 13 else '',
            'leave': row[14] if len(row) > 14 else '',
        }
    return result


def build_matrix(day_numbers: List[int], employees: Dict[str, dict], summary: Dict[str, dict]) -> pd.DataFrame:
    """Build the matrix DataFrame expected by attendance_calculator.py."""
    # Order employees by numeric employee number
    ordered_nos = sorted(employees.keys(), key=lambda x: int(x) if x.isdigit() else float('inf'))

    rows = []
    for emp_no in ordered_nos:
        emp = employees[emp_no]
        info = summary.get(emp_no, {})
        row = {
            '工号': emp_no,
            '姓名': info.get('name', emp['name']),
            '部门': info.get('dept', emp['dept']),
        }
        for day in day_numbers:
            row[f'07/{day:02d}'] = '\n'.join(emp['daily'].get(day, []))
        rows.append(row)

    return pd.DataFrame(rows)


def convert(input_path: str, output_path: Optional[str] = None) -> str:
    """Convert original .xls to matrix .xlsx. Returns output path."""
    path = Path(input_path)
    wb = xlrd.open_workbook(str(path), formatting_info=True)

    day_numbers, employees = extract_employees_from_logs(wb)
    summary = extract_summary(wb)

    df = build_matrix(day_numbers, employees, summary)

    out_path = Path(output_path) if output_path else path.with_stem(path.stem + '_矩阵').with_suffix('.xlsx')

    with pd.ExcelWriter(str(out_path), engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='打卡记录', index=False)

    # Basic formatting
    wb_out = openpyxl.load_workbook(str(out_path))
    ws = wb_out.active
    header_fill = openpyxl.styles.PatternFill('solid', fgColor='404040')
    header_font = openpyxl.styles.Font(color='FFFFFF', bold=True)
    center_align = openpyxl.styles.Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = openpyxl.styles.Side(style='thin', color='BFBFBF')
    border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = border

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value) if cell.value is not None else ''
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 20)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = center_align
            cell.border = border
        ws.row_dimensions[row[0].row].height = 36

    ws.freeze_panes = 'D2'
    wb_out.save(str(out_path))

    return str(out_path)


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description='Convert attendance .xls to matrix .xlsx')
    parser.add_argument('input', help='Input .xls file path')
    parser.add_argument('-o', '--output', default=None, help='Output .xlsx file path')
    args = parser.parse_args(argv)

    if not Path(args.input).exists():
        print(f'Error: file not found: {args.input}', file=sys.stderr)
        return 1

    out = convert(args.input, args.output)
    print(f'Converted matrix saved to: {out}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

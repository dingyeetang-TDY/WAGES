# -*- coding: utf-8 -*-
"""Highlight all odd-punch cells in the matrix Excel for manual correction.

Idempotent: running again on a previously-annotated file will re-evaluate
all date cells and remove/add red highlighting based on current content.
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path
from typing import Optional, Sequence, Tuple

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.comments import Comment


NOTE_TEXT = '标注说明：红色背景单元格 = 奇数打卡（无法配对），请在原数据源中补录/删除一个时间。'


def count_valid_times(cell_text: str) -> int:
    """Count valid HH:MM time strings in a cell."""
    if not cell_text:
        return 0
    count = 0
    for line in str(cell_text).split('\n'):
        line = line.strip()
        if not line:
            continue
        if re.fullmatch(r'\d{1,2}:\d{2}', line):
            try:
                h, m = map(int, line.split(':'))
                if 0 <= h < 24 and 0 <= m < 60:
                    count += 1
            except ValueError:
                pass
    return count


def _is_note_row(ws, row_idx: int) -> bool:
    """Check if the given row is an existing note row."""
    first_cell = ws.cell(row=row_idx, column=1).value
    if first_cell and NOTE_TEXT in str(first_cell):
        return True
    return False


def _find_header_row(ws) -> Tuple[int, bool]:
    """Return (header_row, inserted_note)."""
    if _is_note_row(ws, 1):
        return 2, False
    return 1, True


def _get_date_columns(ws, header_row: int):
    """Return list of (col_index, col_letter) for date columns."""
    date_cols = []
    for col_idx in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=col_idx).value
        if val and re.search(r'\d{2}/\d{2}', str(val)):
            date_cols.append((col_idx, openpyxl.utils.get_column_letter(col_idx)))
    return date_cols


def highlight_odd_punches(input_path: str, output_path: Optional[str] = None) -> Tuple[str, int, int]:
    """Open matrix xlsx, highlight odd-punch cells, save.

    Idempotent: re-running updates styling according to current cell values.
    """
    path = Path(input_path)
    wb = openpyxl.load_workbook(str(path))
    ws = wb.active  # expect single-sheet matrix

    header_row, inserted_note = _find_header_row(ws)

    # If there is no note row yet, insert one
    if inserted_note:
        ws.insert_rows(1)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ws.max_column)
        note_cell = ws.cell(row=1, column=1, value=NOTE_TEXT)
        note_cell.font = Font(color='B71C1C', bold=True, size=11)
        note_cell.fill = PatternFill('solid', fgColor='FFEBEE')
        note_cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 30
        header_row = 2

    date_cols = _get_date_columns(ws, header_row)

    # Style definitions
    odd_fill = PatternFill('solid', fgColor='FFCDD2')       # light red background
    odd_font = Font(color='B71C1C', bold=True)              # dark red bold text
    odd_border = Border(
        left=Side(style='medium', color='E53935'),
        right=Side(style='medium', color='E53935'),
        top=Side(style='medium', color='E53935'),
        bottom=Side(style='medium', color='E53935'),
    )

    default_font = Font(color='000000', bold=False)
    default_border = Border(
        left=Side(style='thin', color='BFBFBF'),
        right=Side(style='thin', color='BFBFBF'),
        top=Side(style='thin', color='BFBFBF'),
        bottom=Side(style='thin', color='BFBFBF'),
    )
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)

    highlighted = 0
    for row_idx in range(header_row + 1, ws.max_row + 1):
        for col_idx, _ in date_cols:
            cell = ws.cell(row=row_idx, column=col_idx)
            cnt = count_valid_times(cell.value)
            if cnt > 0 and cnt % 2 == 1:
                cell.fill = odd_fill
                cell.font = odd_font
                cell.border = odd_border
                cell.alignment = center_align
                cell.comment = Comment(
                    f'奇数打卡：本单元格有 {cnt} 个有效打卡时间，无法配对。\n'
                    '请补录或删除一个时间，使其成为偶数。',
                    'Attendance Calculator'
                )
                highlighted += 1
            else:
                # Clear previous odd-punch styling if it was set
                cell.fill = PatternFill(fill_type=None)
                cell.font = default_font
                cell.border = default_border
                cell.alignment = center_align
                cell.comment = None

    # Always re-apply header formatting
    header_fill = PatternFill('solid', fgColor='404040')
    header_font = Font(color='FFFFFF', bold=True)
    for cell in ws[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = default_border

    ws.freeze_panes = f'D{header_row + 1}'

    out_path = Path(output_path) if output_path else path.with_stem(path.stem + '_奇数标注').with_suffix('.xlsx')
    wb.save(str(out_path))
    return str(out_path), highlighted, len(date_cols)


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description='Highlight odd-punch cells in attendance matrix xlsx')
    parser.add_argument('input', help='Input matrix .xlsx file path')
    parser.add_argument('-o', '--output', default=None, help='Output .xlsx file path')
    args = parser.parse_args(argv)

    if not Path(args.input).exists():
        print(f'Error: file not found: {args.input}', file=sys.stderr)
        return 1

    out, highlighted, date_count = highlight_odd_punches(args.input, args.output)
    print(f'Output saved to: {out}')
    print(f'Date columns scanned: {date_count}')
    print(f'Odd-punch cells highlighted: {highlighted}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

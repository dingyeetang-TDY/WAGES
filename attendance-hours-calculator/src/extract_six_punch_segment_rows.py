# -*- coding: utf-8 -*-
"""Extract 6-punch employees and display each work/break segment as independent rows.

Output format per employee sheet:
  A: DAY
  B: SEGMENT
  C: TYPE (WORK / BREAK 1 / BREAK 2)
  D: START TIME
  E: END TIME
  F: DURATION
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
from pathlib import Path
from typing import List, Optional, Sequence

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd


def parse_date_columns(columns) -> List[str]:
    return [str(c).strip() for c in columns if re.match(r'\d{2}/\d{2}', str(c).strip())]


def parse_punch_times(cell_value) -> List[datetime.time]:
    if cell_value is None or (isinstance(cell_value, float) and pd.isna(cell_value)):
        return []
    parts = re.split(r'[\n,;/]+', str(cell_value))
    times = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        for fmt in ('%H:%M', '%H:%M:%S'):
            try:
                times.append(datetime.datetime.strptime(part, fmt).time())
                break
            except ValueError:
                continue
    return sorted(times)


def minutes(t: Optional[datetime.time]) -> float:
    if t is None:
        return 0.0
    return t.hour * 60 + t.minute + t.second / 60.0


def timedelta_from_minutes(mins: float) -> datetime.timedelta:
    return datetime.timedelta(minutes=mins)


def has_six_punch_day(daily_data: dict) -> bool:
    return any(len(times) >= 6 for times in daily_data.values())


def create_employee_sheet(wb, employee_name: str, employee_no, dept: str, daily_data: dict):
    safe_name = re.sub(r'[\\/*?:\[\]]', '-', str(employee_name))[:31]
    ws = wb.create_sheet(title=safe_name)

    ws['A1'] = f'{employee_name} ({employee_no}) - {dept}'
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:F1')

    headers = ['DAY', 'SEGMENT', 'TYPE', 'START TIME', 'END TIME', 'DURATION']
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDDDDD')
        cell.alignment = Alignment(horizontal='center')

    days = sorted(daily_data.keys(), key=lambda d: int(d.split('/')[1]))
    row_idx = 4
    total_work = 0.0
    total_break = 0.0

    for day_str in days:
        day_num = int(day_str.split('/')[1])
        times = daily_data[day_str]

        if not times:
            continue

        if len(times) % 2 == 1:
            # Odd punch - just list as error row
            ws.cell(row=row_idx, column=1, value=day_num)
            ws.cell(row=row_idx, column=2, value='-')
            ws.cell(row=row_idx, column=3, value='ODD PUNCH')
            ws.cell(row=row_idx, column=4, value=times[0])
            ws.cell(row=row_idx, column=5, value=times[-1])
            for col in range(1, 7):
                ws.cell(row=row_idx, column=col).fill = PatternFill('solid', fgColor='FFC7CE')
            row_idx += 1
            continue

        for i in range(len(times) - 1):
            start_t = times[i]
            end_t = times[i + 1]
            duration = minutes(end_t) - minutes(start_t)

            if i % 2 == 0:
                seg_type = 'WORK'
                segment_no = (i // 2) + 1
                total_work += duration
            else:
                seg_type = f'BREAK {(i // 2) + 1}'
                segment_no = (i // 2) + 1
                total_break += duration

            ws.cell(row=row_idx, column=1, value=day_num)
            ws.cell(row=row_idx, column=2, value=segment_no)
            ws.cell(row=row_idx, column=3, value=seg_type)
            ws.cell(row=row_idx, column=4, value=start_t)
            ws.cell(row=row_idx, column=5, value=end_t)
            ws.cell(row=row_idx, column=6, value=timedelta_from_minutes(duration))
            row_idx += 1

    summary_row = row_idx + 1
    ws.cell(row=summary_row, column=1, value='TOTAL WORK')
    ws.cell(row=summary_row, column=6, value=timedelta_from_minutes(total_work))
    ws.cell(row=summary_row + 1, column=1, value='TOTAL BREAK')
    ws.cell(row=summary_row + 1, column=6, value=timedelta_from_minutes(total_break))

    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for r in range(1, ws.max_row + 1):
            cell = ws.cell(row=r, column=col_idx)
            try:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 20)


def extract_six_punch_employees(input_path: str, output_path: str):
    df = pd.read_excel(input_path, header=1)
    date_cols = parse_date_columns(df.columns)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    extracted = 0
    for idx, row in df.iterrows():
        employee_no = row.get('工号', idx + 1)
        name = row.get('姓名', f'Employee_{idx + 1}')
        dept = row.get('部门', '')
        daily_data = {}
        for col in date_cols:
            daily_data[col] = parse_punch_times(row.get(col))

        if has_six_punch_day(daily_data):
            create_employee_sheet(wb, name, employee_no, dept, daily_data)
            extracted += 1

    wb.save(output_path)
    print(f'Extracted {extracted} employees with 6-punch days.')
    print(f'Output saved: {output_path}')


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description='Extract 6-punch employees with segment rows')
    parser.add_argument('input', help='Input matrix .xlsx')
    parser.add_argument('-o', '--output', default='6次打卡员工_分段明细.xlsx', help='Output path')
    args = parser.parse_args(argv)

    if not Path(args.input).exists():
        print(f'Error: file not found: {args.input}', file=sys.stderr)
        return 1

    extract_six_punch_employees(args.input, args.output)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

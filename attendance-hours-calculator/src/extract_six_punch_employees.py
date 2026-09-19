# -*- coding: utf-8 -*-
"""Extract employees with 6-punch days and show each break period independently.

Input: attendance matrix .xlsx
Output: .xlsx with one sheet per 6-punch employee, columns:
  A: DAY
  B: START TIME
  C: END TIME
  D: WORKING HOURS (gross)
  E-G: BREAK 1 START / END / DURATION
  H-J: BREAK 2 START / END / DURATION
  K: NET WORKING HOURS
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
from pathlib import Path
from typing import List, Optional, Sequence

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd


def parse_date_columns(columns) -> List[str]:
    return [str(c).strip() for c in columns if re.match(r'\d{2}/\d{2}', str(c).strip())]


def parse_punch_times(cell_value) -> List[datetime.time]:
    if cell_value is None or (isinstance(cell_value, float) and pd.isna(cell_value)):
        return []
    parts = re.split(r'[\n,;/]+', str(cell_value))
    times = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        for fmt in ('%H:%M', '%H:%M:%S'):
            try:
                times.append(datetime.datetime.strptime(part, fmt).time())
                break
            except ValueError:
                continue
    return sorted(times)


def minutes(t: Optional[datetime.time]) -> float:
    if t is None:
        return 0.0
    return t.hour * 60 + t.minute + t.second / 60.0


def timedelta_from_minutes(mins: float) -> datetime.timedelta:
    return datetime.timedelta(minutes=mins)


def has_six_punch_day(daily_data: dict) -> bool:
    return any(len(times) >= 6 for times in daily_data.values())


def compute_day_with_breaks(times: List[datetime.time]):
    """Compute work/break segments. Return dict with all details."""
    result = {
        'start': None, 'end': None,
        'gross': 0.0, 'net': 0.0,
        'breaks': [],  # list of {'start': time, 'end': time, 'duration': minutes}
        'is_odd': len(times) % 2 == 1,
    }
    if not times:
        return result

    result['start'] = times[0]
    result['end'] = times[-1]
    result['gross'] = minutes(times[-1]) - minutes(times[0])

    if result['is_odd']:
        return result

    for i in range(len(times) - 1):
        segment = minutes(times[i + 1]) - minutes(times[i])
        if i % 2 == 0:
            result['net'] += segment
        else:
            result['breaks'].append({
                'start': times[i],
                'end': times[i + 1],
                'duration': segment,
            })

    return result


def create_employee_sheet(wb, employee_name: str, employee_no, dept: str, daily_data: dict):
    safe_name = re.sub(r'[\\/*?:\[\]]', '-', str(employee_name))[:31]
    ws = wb.create_sheet(title=safe_name)

    ws['A1'] = f'{employee_name} ({employee_no}) - {dept}'
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:K1')

    headers = [
        'DAY', 'START TIME', 'END TIME', 'WORKING HOURS',
        'BREAK 1 START', 'BREAK 1 END', 'BREAK 1 DURATION',
        'BREAK 2 START', 'BREAK 2 END', 'BREAK 2 DURATION',
        'NET WORKING HOURS'
    ]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDDDDD')
        cell.alignment = Alignment(horizontal='center')

    total_gross = 0.0
    total_net = 0.0
    total_breaks = [0.0, 0.0]
    work_days = 0

    days = sorted(daily_data.keys(), key=lambda d: int(d.split('/')[1]))
    for row_idx, day_str in enumerate(days, start=4):
        day_num = int(day_str.split('/')[1])
        times = daily_data[day_str]
        day_data = compute_day_with_breaks(times)

        ws.cell(row=row_idx, column=1, value=day_num)
        if day_data['start']:
            ws.cell(row=row_idx, column=2, value=day_data['start'])
        if day_data['end']:
            ws.cell(row=row_idx, column=3, value=day_data['end'])
        if day_data['gross'] > 0:
            ws.cell(row=row_idx, column=4, value=timedelta_from_minutes(day_data['gross']))

        for b_idx, break_info in enumerate(day_data['breaks'][:2]):
            col_offset = 5 + b_idx * 3
            ws.cell(row=row_idx, column=col_offset, value=break_info['start'])
            ws.cell(row=row_idx, column=col_offset + 1, value=break_info['end'])
            ws.cell(row=row_idx, column=col_offset + 2, value=timedelta_from_minutes(break_info['duration']))
            total_breaks[b_idx] += break_info['duration']

        if day_data['net'] > 0:
            ws.cell(row=row_idx, column=11, value=timedelta_from_minutes(day_data['net']))
            total_gross += day_data['gross']
            total_net += day_data['net']
            work_days += 1

        if day_data['is_odd']:
            for col in range(1, 12):
                ws.cell(row=row_idx, column=col).fill = PatternFill('solid', fgColor='FFC7CE')

    summary_row = row_idx + 2
    ws.cell(row=summary_row, column=1, value='TOTAL WORKING HOURS')
    ws.cell(row=summary_row, column=4, value=timedelta_from_minutes(total_gross))
    ws.cell(row=summary_row, column=11, value=timedelta_from_minutes(total_net))

    ws.cell(row=summary_row + 1, column=1, value='TOTAL BREAK TIME')
    for b_idx in range(2):
        col_offset = 5 + b_idx * 3
        ws.cell(row=summary_row + 1, column=col_offset + 2, value=timedelta_from_minutes(total_breaks[b_idx]))

    ws.cell(row=summary_row + 2, column=1, value=f'WORK DAYS: {work_days}')

    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for r in range(1, ws.max_row + 1):
            cell = ws.cell(row=r, column=col_idx)
            try:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 20)


def extract_six_punch_employees(input_path: str, output_path: str):
    df = pd.read_excel(input_path, header=1)
    date_cols = parse_date_columns(df.columns)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    extracted = 0
    for idx, row in df.iterrows():
        employee_no = row.get('工号', idx + 1)
        name = row.get('姓名', f'Employee_{idx + 1}')
        dept = row.get('部门', '')
        daily_data = {}
        for col in date_cols:
            daily_data[col] = parse_punch_times(row.get(col))

        if has_six_punch_day(daily_data):
            create_employee_sheet(wb, name, employee_no, dept, daily_data)
            extracted += 1

    wb.save(output_path)
    print(f'Extracted {extracted} employees with 6-punch days.')
    print(f'Output saved: {output_path}')


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description='Extract 6-punch employees with independent break columns')
    parser.add_argument('input', help='Input matrix .xlsx')
    parser.add_argument('-o', '--output', default='6次打卡员工_独立休息段.xlsx', help='Output path')
    args = parser.parse_args(argv)

    if not Path(args.input).exists():
        print(f'Error: file not found: {args.input}', file=sys.stderr)
        return 1

    extract_six_punch_employees(args.input, args.output)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

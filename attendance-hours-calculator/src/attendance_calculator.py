# -*- coding: utf-8 -*-
"""Attendance / work-hours calculator for Excel punch logs.

Input format
------------
An .xlsx file with columns such as:
    工号 | 姓名 | 部门 | 07/01 | 07/02 | ... | 07/18
Each date cell contains newline-separated punch times, e.g.:
    08:50
    18:46

Calculation rules
-----------------
1. Parse every "HH:MM" time, sort them chronologically.
2. Pair consecutive punches:
   - pairs 1-2, 3-4, 5-6, ... are WORK segments (counted as work hours)
   - gaps 2-3, 4-5, ... are BREAK segments (mid-break, not counted)
3. If a day has an odd number of valid punches, mark it as anomaly;
   work hours and break hours for that day are 0.
4. Late / early leave are measured against configurable standard times
   (default AM in = 08:30, PM out = 17:30).

Output
------
An .xlsx workbook with four sheets:
    打卡记录   - original punch times
    每日工时   - work hours per employee per day
    休息时长   - break duration per employee per day
    汇总统计   - summary statistics
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Tuple

import openpyxl
import pandas as pd


# ---------------------------------------------------------------------------
# Time helpers
# ---------------------------------------------------------------------------

def parse_time(time_str: str) -> Optional[int]:
    """Convert 'HH:MM' (or 'H:MM') to minutes since midnight."""
    if not time_str or not isinstance(time_str, str):
        return None
    time_str = time_str.strip()
    if not time_str:
        return None
    m = re.fullmatch(r'(\d{1,2}):(\d{2})', time_str)
    if not m:
        return None
    h, mi = int(m.group(1)), int(m.group(2))
    if not (0 <= h < 24 and 0 <= mi < 60):
        return None
    return h * 60 + mi


def format_minutes(minutes: Optional[int]) -> str:
    """Minutes since midnight -> 'HH:MM'."""
    if minutes is None:
        return ''
    return f'{minutes // 60:02d}:{minutes % 60:02d}'


def hours_from_minutes(minutes: int) -> float:
    """Return hours rounded to 2 decimals."""
    return round(minutes / 60.0, 2)


# ---------------------------------------------------------------------------
# Core calculation
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class Segment:
    start: int  # minutes since midnight
    end: int

    @property
    def minutes(self) -> int:
        return self.end - self.start

    def __repr__(self) -> str:
        return f'{format_minutes(self.start)}-{format_minutes(self.end)}'


@dataclass
class DayResult:
    raw_punches: List[str]
    valid_count: int
    is_odd: bool
    work_segments: List[Segment] = field(default_factory=list)
    break_segments: List[Segment] = field(default_factory=list)

    @property
    def work_minutes(self) -> int:
        return sum(s.minutes for s in self.work_segments)

    @property
    def break_minutes(self) -> int:
        return sum(s.minutes for s in self.break_segments)

    @property
    def total_span_minutes(self) -> int:
        if self.valid_count == 0:
            return 0
        minutes = sorted(m for m in (parse_time(t) for t in self.raw_punches) if m is not None)
        return minutes[-1] - minutes[0]

    @property
    def first_punch_minutes(self) -> Optional[int]:
        minutes = sorted(m for m in (parse_time(t) for t in self.raw_punches) if m is not None)
        return minutes[0] if minutes else None

    @property
    def last_punch_minutes(self) -> Optional[int]:
        minutes = sorted(m for m in (parse_time(t) for t in self.raw_punches) if m is not None)
        return minutes[-1] if minutes else None


def split_punches(raw_text: str) -> List[str]:
    """Split a cell value into individual time strings."""
    if not raw_text:
        return []
    return [part.strip() for part in str(raw_text).split('\n') if part.strip()]


def calculate_day_hours(punches: Iterable[str]) -> dict:
    """Apply the work/break pairing rules to one day's punches.

    Returns a dict with keys:
        raw_punches, count, is_odd, work_segments, break_segments,
        work_minutes, break_minutes, total_span_minutes,
        first_punch_minutes, last_punch_minutes.
    Segments are returned as (start_min, end_min, duration_min) tuples.
    """
    raw = list(punches)
    minutes = sorted(m for m in (parse_time(t) for t in raw) if m is not None)
    count = len(minutes)
    is_odd = (count % 2 == 1)

    result = {
        'raw_punches': raw,
        'count': count,
        'is_odd': is_odd,
        'work_segments': [],
        'break_segments': [],
    }

    if count == 0 or is_odd:
        result['work_minutes'] = 0
        result['break_minutes'] = 0
        result['total_span_minutes'] = 0
        result['first_punch_minutes'] = None
        result['last_punch_minutes'] = None
        return result

    # work segments: (0,1), (2,3), ...
    work_segments = []
    for i in range(0, count, 2):
        work_segments.append((minutes[i], minutes[i + 1], minutes[i + 1] - minutes[i]))

    # break segments: (1,2), (3,4), ...
    break_segments = []
    for i in range(1, count - 1, 2):
        break_segments.append((minutes[i], minutes[i + 1], minutes[i + 1] - minutes[i]))

    result['work_segments'] = work_segments
    result['break_segments'] = break_segments
    result['work_minutes'] = sum(s[2] for s in work_segments)
    result['break_minutes'] = sum(s[2] for s in break_segments)
    result['total_span_minutes'] = minutes[-1] - minutes[0]
    result['first_punch_minutes'] = minutes[0]
    result['last_punch_minutes'] = minutes[-1]
    return result


# ---------------------------------------------------------------------------
# Excel I/O
# ---------------------------------------------------------------------------

def identify_date_columns(df: pd.DataFrame, date_regex: str = r'\d{2}/\d{2}') -> List[str]:
    """Return column names that look like dates (default MM/DD)."""
    return [c for c in df.columns if re.search(date_regex, str(c))]


def read_punch_excel(path: str, header_row: int = 1) -> pd.DataFrame:
    """Read the source Excel (.xlsx or .xls) and keep all columns as string/object.

    header_row is 1-based. Default = 1 (first row is header).
    Use header_row=4 if the first three rows are titles/notes.
    """
    path_lower = path.lower()
    engine = 'xlrd' if path_lower.endswith('.xls') else None
    header_idx = header_row - 1  # pandas uses 0-based index
    if engine:
        return pd.read_excel(path, dtype=str, engine='xlrd', header=header_idx)
    return pd.read_excel(path, dtype=str, header=header_idx)


def is_standard_weekend(date_label: str) -> bool:
    """Best-effort weekend detection from a column label like '07/01 We'."""
    label = str(date_label).strip().split()[-1]  # last token
    return label in ('Sa', 'Su', 'Sat', 'Sun', '六', '日')


# ---------------------------------------------------------------------------
# Reporting
# ---------------------------------------------------------------------------

def process_attendance(
    df: pd.DataFrame,
    id_col: str = '工号',
    name_col: str = '姓名',
    dept_col: str = '部门',
    date_cols: Optional[Sequence[str]] = None,
    am_in: str = '08:30',
    pm_out: str = '17:30',
) -> dict:
    """Calculate work/break hours for every employee and every date column."""
    if date_cols is None:
        date_cols = identify_date_columns(df)

    am_min = parse_time(am_in) or 510
    pm_min = parse_time(pm_out) or 1050

    records = []
    for _, row in df.iterrows():
        emp_id = row.get(id_col, '')
        emp_name = row.get(name_col, '')
        emp_dept = row.get(dept_col, '')
        record = {
            id_col: emp_id,
            name_col: emp_name,
            dept_col: emp_dept,
            '_days': {},
        }
        for col in date_cols:
            punches = split_punches(row.get(col))
            day = calculate_day_hours(punches)
            record['_days'][col] = day
        records.append(record)

    # Build per-metric DataFrames
    work_rows, break_rows, raw_rows, summary_rows = [], [], [], []

    for rec in records:
        emp_id, emp_name, emp_dept = rec[id_col], rec[name_col], rec[dept_col]
        raw_row = {id_col: emp_id, name_col: emp_name, dept_col: emp_dept}
        work_row = dict(raw_row)
        break_row = dict(raw_row)

        total_work_min = 0
        total_break_min = 0
        total_span_min = 0
        present_days = 0
        anomaly_days = 0
        late_count = 0
        late_minutes = 0
        early_count = 0
        early_minutes = 0
        first_punch: Optional[int] = None
        last_punch: Optional[int] = None

        for col in date_cols:
            day = rec['_days'][col]
            raw_row[col] = '\n'.join(day['raw_punches'])

            if day['count'] == 0:
                work_row[col] = ''
                break_row[col] = ''
                continue

            present_days += 1
            if day['is_odd']:
                anomaly_days += 1
                work_row[col] = 0.0
                break_row[col] = 0.0
                continue

            work_min = day['work_minutes']
            break_min = day['break_minutes']
            total_work_min += work_min
            total_break_min += break_min
            total_span_min += day['total_span_minutes']
            work_row[col] = hours_from_minutes(work_min)
            break_row[col] = hours_from_minutes(break_min)

            if day['first_punch_minutes'] is not None:
                if first_punch is None or day['first_punch_minutes'] < first_punch:
                    first_punch = day['first_punch_minutes']
                if last_punch is None or day['last_punch_minutes'] > last_punch:
                    last_punch = day['last_punch_minutes']
                if day['first_punch_minutes'] > am_min:
                    late_count += 1
                    late_minutes += day['first_punch_minutes'] - am_min
                if day['last_punch_minutes'] < pm_min:
                    early_count += 1
                    early_minutes += pm_min - day['last_punch_minutes']

        raw_rows.append(raw_row)
        work_rows.append(work_row)
        break_rows.append(break_row)

        summary_rows.append({
            id_col: emp_id,
            name_col: emp_name,
            dept_col: emp_dept,
            '出勤天数': present_days,
            '打卡总天数': present_days,
            '奇数打卡天数': anomaly_days,
            '工时(工作时段,h)': hours_from_minutes(total_work_min),
            '休息时长(h)': hours_from_minutes(total_break_min),
            '在司总时长(h)': hours_from_minutes(total_span_min),
            '首次打卡': format_minutes(first_punch),
            '末次打卡': format_minutes(last_punch),
            '迟到次数': late_count,
            '迟到分钟': late_minutes,
            '早退次数': early_count,
            '早退分钟': early_minutes,
        })

    return {
        'raw': pd.DataFrame(raw_rows),
        'work': pd.DataFrame(work_rows),
        'break': pd.DataFrame(break_rows),
        'summary': pd.DataFrame(summary_rows),
        'date_columns': list(date_cols),
    }


def write_results(results: dict, output_path: str) -> None:
    """Write the four result DataFrames to one Excel workbook."""
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        results['raw'].to_excel(writer, sheet_name='打卡记录', index=False)
        results['work'].to_excel(writer, sheet_name='每日工时', index=False)
        results['break'].to_excel(writer, sheet_name='休息时长', index=False)
        results['summary'].to_excel(writer, sheet_name='汇总统计', index=False)

    # Apply some basic formatting
    wb = openpyxl.load_workbook(output_path)
    header_fill = openpyxl.styles.PatternFill('solid', fgColor='404040')
    header_font = openpyxl.styles.Font(color='FFFFFF', bold=True)
    center_align = openpyxl.styles.Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = openpyxl.styles.Side(style='thin', color='BFBFBF')
    border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = border
        # Auto width
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    val = str(cell.value) if cell.value is not None else ''
                    max_len = max(max_len, len(val))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 2, 30)
        # Freeze header + first three columns
        ws.freeze_panes = 'D2'

    wb.save(output_path)


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def build_arg_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description='Calculate work hours and break duration from attendance Excel.'
    )
    parser.add_argument('input', help='Input .xlsx file path')
    parser.add_argument('-o', '--output', default=None, help='Output .xlsx file path')
    parser.add_argument('--id-col', default='工号', help='Employee ID column name')
    parser.add_argument('--name-col', default='姓名', help='Employee name column name')
    parser.add_argument('--dept-col', default='部门', help='Department column name')
    parser.add_argument('--am-in', default='08:30', help='Standard morning start time HH:MM')
    parser.add_argument('--pm-out', default='17:30', help='Standard evening end time HH:MM')
    parser.add_argument('--header-row', type=int, default=1,
                        help='Row number containing column headers (1-based; default 1)')
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_arg_parser()
    args = parser.parse_args(argv)

    input_path = Path(args.input)
    if not input_path.exists():
        print(f'Error: file not found: {input_path}', file=sys.stderr)
        return 1

    output_path = args.output or input_path.with_stem(input_path.stem + '_工时计算结果').with_suffix('.xlsx')

    df = read_punch_excel(str(input_path), header_row=args.header_row)
    date_cols = identify_date_columns(df)
    if not date_cols:
        print('Warning: no date-like columns found (expected e.g. 07/01).', file=sys.stderr)

    results = process_attendance(
        df,
        id_col=args.id_col,
        name_col=args.name_col,
        dept_col=args.dept_col,
        date_cols=date_cols,
        am_in=args.am_in,
        pm_out=args.pm_out,
    )

    write_results(results, str(output_path))
    print(f'Results saved to: {output_path}')
    print(f'Employees: {len(results["summary"])}, Date columns: {len(date_cols)}')
    print(f'Total work hours: {results["summary"]["工时(工作时段,h)"].sum():.2f}')
    print(f'Total break hours: {results["summary"]["休息时长(h)"].sum():.2f}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

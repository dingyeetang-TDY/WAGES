# -*- coding: utf-8 -*-
"""Convert attendance matrix to JUN.xlsx per-employee format.

Input: matrix .xlsx from convert_xls.py / pipeline
Output: .xlsx with one sheet per employee, columns similar to JUN.xlsx:
  A: DAY
  B: START TIME
  C: END TIME
  D: WORKING HOURS (gross in-company time)
  E: BREAK TIME (rest between punch pairs)
  F: NET WORKING HOURS (working periods only)
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd


def parse_date_columns(columns) -> List[str]:
    """Return column names that look like MM/DD dates."""
    date_cols = []
    for c in columns:
        s = str(c).strip()
        if re.match(r'\d{2}/\d{2}', s):
            date_cols.append(s)
    return date_cols


def parse_punch_times(cell_value) -> List[datetime.time]:
    """Parse one or more HH:MM times from a cell."""
    if cell_value is None or (isinstance(cell_value, float) and pd.isna(cell_value)):
        return []
    text = str(cell_value)
    # Split on newline, comma, or slash
    parts = re.split(r'[\n,;/]+', text)
    times = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        try:
            t = datetime.datetime.strptime(part, '%H:%M').time()
            times.append(t)
        except ValueError:
            try:
                t = datetime.datetime.strptime(part, '%H:%M:%S').time()
                times.append(t)
            except ValueError:
                continue
    return sorted(times)


def minutes(t: datetime.time) -> float:
    return t.hour * 60 + t.minute + t.second / 60.0


def time_from_minutes(mins: float) -> datetime.time:
    mins = max(0, mins)
    h = int(mins // 60) % 24
    m = int(mins % 60)
    s = int(round((mins - int(mins)) * 60)) % 60
    return datetime.time(h, m, s)


def timedelta_from_minutes(mins: float) -> datetime.timedelta:
    return datetime.timedelta(minutes=mins)


def compute_day(times: List[datetime.time]) -> Tuple[Optional[datetime.time], Optional[datetime.time], float, float, float, bool]:
    """Compute start, end, gross, break, net for a list of punch times.

    Pairing rule:
      - (1-2), (3-4), (5-6)... -> working periods
      - (2-3), (4-5)... -> break/rest periods
      - Odd number of punches -> error
    """
    if not times:
        return None, None, 0.0, 0.0, 0.0, False

    start = times[0]
    end = times[-1]
    gross = minutes(end) - minutes(start)

    if len(times) % 2 == 1:
        return start, end, gross, 0.0, 0.0, True

    work = 0.0
    break_mins = 0.0
    for i in range(len(times) - 1):
        segment = minutes(times[i + 1]) - minutes(times[i])
        # Pair index: i=0 (1-2) is work, i=1 (2-3) is break, i=2 (3-4) is work...
        if i % 2 == 0:
            work += segment
        else:
            break_mins += segment

    return start, end, gross, break_mins, work, False


def create_employee_sheet(wb, employee_name: str, employee_no, dept: str, daily_data: dict):
    """Create one sheet in JUN format for an employee."""
    # Sanitize sheet name (Excel limit 31 chars, no special chars)
    safe_name = re.sub(r'[\\/*?:\[\]]', '-', str(employee_name))[:31]
    ws = wb.create_sheet(title=safe_name)

    # Title
    ws['A1'] = f'{employee_name} ({employee_no}) - {dept}'
    ws['A1'].font = Font(bold=True, size=14)
    ws.merge_cells('A1:F1')

    # Header row
    headers = ['DAY', 'START TIME', 'END TIME', 'WORKING HOURS', 'BREAK TIME', 'NET WORKING HOURS']
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDDDDD')
        cell.alignment = Alignment(horizontal='center')

    # Daily rows
    total_gross = 0.0
    total_break = 0.0
    total_net = 0.0
    work_days = 0

    days = sorted(daily_data.keys(), key=lambda d: int(d.split('/')[1]))
    for row_idx, day_str in enumerate(days, start=4):
        day_num = int(day_str.split('/')[1])
        times = daily_data[day_str]
        start, end, gross, break_mins, net, is_odd = compute_day(times)

        ws.cell(row=row_idx, column=1, value=day_num)
        if start:
            ws.cell(row=row_idx, column=2, value=start)
        if end:
            ws.cell(row=row_idx, column=3, value=end)
        if gross > 0:
            ws.cell(row=row_idx, column=4, value=timedelta_from_minutes(gross))
        if break_mins > 0:
            ws.cell(row=row_idx, column=5, value=timedelta_from_minutes(break_mins))
        if net > 0:
            ws.cell(row=row_idx, column=6, value=timedelta_from_minutes(net))
            total_gross += gross
            total_break += break_mins
            total_net += net
            work_days += 1
        if is_odd:
            for col in range(1, 7):
                ws.cell(row=row_idx, column=col).fill = PatternFill('solid', fgColor='FFC7CE')

    # Summary rows
    summary_row = row_idx + 2
    ws.cell(row=summary_row, column=1, value='TOTAL WORKING HOURS')
    ws.cell(row=summary_row, column=4, value=timedelta_from_minutes(total_gross))
    ws.cell(row=summary_row, column=6, value=timedelta_from_minutes(total_net))

    ws.cell(row=summary_row + 1, column=1, value='TOTAL BREAK TIME')
    ws.cell(row=summary_row + 1, column=5, value=timedelta_from_minutes(total_break))

    ws.cell(row=summary_row + 2, column=1, value=f'WORK DAYS: {work_days}')

    # Format columns
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            try:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 20)


def matrix_to_jun_format(input_path: str, output_path: str):
    df = pd.read_excel(input_path, header=1)
    date_cols = parse_date_columns(df.columns)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # remove default sheet

    for idx, row in df.iterrows():
        employee_no = row.get('工号', idx + 1)
        name = row.get('姓名', f'Employee_{idx + 1}')
        dept = row.get('部门', '')
        daily_data = {}
        for col in date_cols:
            daily_data[col] = parse_punch_times(row.get(col))
        create_employee_sheet(wb, name, employee_no, dept, daily_data)

    wb.save(output_path)
    print(f'JUN-format workbook saved: {output_path}')
    print(f'Total employees: {len(df)}')
    print(f'Date columns: {len(date_cols)}')


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = argparse.ArgumentParser(description='Convert attendance matrix to JUN-format per-employee sheets')
    parser.add_argument('input', help='Input matrix .xlsx')
    parser.add_argument('-o', '--output', default='JUN格式_工时统计.xlsx', help='Output path')
    args = parser.parse_args(argv)

    if not Path(args.input).exists():
        print(f'Error: file not found: {args.input}', file=sys.stderr)
        return 1

    matrix_to_jun_format(args.input, args.output)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())

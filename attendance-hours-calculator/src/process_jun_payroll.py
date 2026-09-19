# -*- coding: utf-8 -*-
"""Process JUN.xlsx (or similar manual salary sheets) through the payroll engine.

Assumptions (from default config):
  - Full-time standard: 7.5h/day, 26 days/month unless parsed from sheet
  - Part-time public holidays are double-counted in final work hours
  - Break deduction: only excess over 60 min standard is deducted
  - Multiple rest segments: one standard allowance per day
  - Public holiday: 2026-06-01 for Malaysia
"""
from __future__ import annotations

import argparse
import datetime
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from payroll_engine import PayrollEngine, DailyRecord
from verify_salary_xlsx import (
    find_col_by_header,
    parse_basic_hours_text,
    find_cell_by_keyword,
)


def time_to_minutes(value) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, datetime.time):
        return value.hour * 60 + value.minute + value.second / 60.0
    if isinstance(value, datetime.timedelta):
        return value.total_seconds() / 60.0
    return None


def td_to_minutes(value) -> Optional[float]:
    if isinstance(value, datetime.timedelta):
        return value.total_seconds() / 60.0
    return None


def find_simple_break_column(ws, working_col: Optional[int]) -> Optional[int]:
    """Find a column containing break durations as datetime.time (e.g., 0:30)."""
    if not working_col:
        return None
    for col_idx in range(working_col + 1, ws.max_column + 1):
        valid = 0
        for row_idx in range(2, min(ws.max_row, 33)):
            val = ws.cell(row=row_idx, column=col_idx).value
            if isinstance(val, datetime.time) and val.hour == 0:
                # Likely a break duration like 0:30
                valid += 1
        if valid >= 3:
            return col_idx
    return None


def find_all_rest_segments(ws) -> List[Tuple[int, int, int]]:
    """Find all (rest_col, back_col, duration_col) pairs in sheet."""
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    segments = []
    # Collect positions of REST TIME and BACK TIME headers
    rest_positions = []
    back_positions = []
    for col_idx, h in enumerate(header, start=1):
        if h and 'REST TIME' in str(h).upper():
            rest_positions.append(col_idx)
        elif h and 'BACK TIME' in str(h).upper():
            back_positions.append(col_idx)

    for i, rest_col in enumerate(rest_positions):
        back_col = back_positions[i] if i < len(back_positions) else None
        if not back_col:
            continue
        # Duration column is between REST and BACK (if any)
        duration_col = None
        for col_idx in range(min(rest_col, back_col) + 1, max(rest_col, back_col)):
            h = header[col_idx - 1]
            if h is None:
                duration_col = col_idx
                break
        segments.append((rest_col, back_col, duration_col))
    return segments


def extract_public_holidays_from_sheet(ws, year: int, month: int) -> set:
    """Scan sheet for 'PUBLIC HOLIDAYS (DD/MM/YYYY)' and return set of day numbers."""
    ph_days = set()
    from verify_salary_xlsx import parse_public_holiday_text, find_cell_by_keyword
    ph_cell = find_cell_by_keyword(ws, 'PUBLIC HOLIDAYS')
    if ph_cell:
        parsed = parse_public_holiday_text(str(ph_cell[2]))
        if parsed:
            day, m, y = parsed
            if m == month and y == year:
                ph_days.add(day)
    return ph_days


def extract_records(ws, year: int, month: int, engine: PayrollEngine) -> List[DailyRecord]:
    """Extract daily records from one employee sheet."""
    records = []
    cfg_rest = engine.config['rest']
    standard_rest = cfg_rest['standard_minutes_per_day']

    ph_days = extract_public_holidays_from_sheet(ws, year, month)

    start_col = find_col_by_header(ws, 'START')
    end_col = find_col_by_header(ws, 'END')
    working_col = find_col_by_header(ws, 'WORKING')
    standard_rest_col = find_col_by_header(ws, 'STANDARD REST') or find_col_by_header(ws, 'STANDARD TIME')
    exceed_rest_col = find_col_by_header(ws, 'EXCEED REST')
    rest_segments = find_all_rest_segments(ws)
    simple_break_col = find_simple_break_column(ws, working_col)

    for row_idx in range(2, ws.max_row + 1):
        day = ws.cell(row=row_idx, column=1).value
        if not isinstance(day, int) or day < 1 or day > 31:
            continue

        rec = DailyRecord(day=day)
        rec.start_minutes = time_to_minutes(ws.cell(row=row_idx, column=start_col).value) if start_col else None
        rec.end_minutes = time_to_minutes(ws.cell(row=row_idx, column=end_col).value) if end_col else None

        # Gross
        gross = None
        if working_col:
            gross = td_to_minutes(ws.cell(row=row_idx, column=working_col).value)
        if gross is None and rec.start_minutes is not None and rec.end_minutes is not None:
            gross = rec.end_minutes - rec.start_minutes
        if gross is not None and gross > 0:
            rec.gross_minutes = gross

        # Public holiday flag (use explicit dates from sheet, not full national calendar)
        rec.is_public_holiday = day in ph_days

        # Actual break and deductible break
        actual_break = 0.0
        if rest_segments:
            for rest_col, back_col, duration_col in rest_segments:
                if duration_col:
                    dur = td_to_minutes(ws.cell(row=row_idx, column=duration_col).value)
                    if dur:
                        actual_break += dur
                else:
                    rest_start = time_to_minutes(ws.cell(row=row_idx, column=rest_col).value)
                    back_time = time_to_minutes(ws.cell(row=row_idx, column=back_col).value)
                    if rest_start is not None and back_time is not None:
                        actual_break += back_time - rest_start
        elif simple_break_col:
            # Simple break column with datetime.time like 0:30
            # When no standard-rest column exists, the whole break is typically deducted
            break_mins = time_to_minutes(ws.cell(row=row_idx, column=simple_break_col).value)
            if break_mins is not None:
                actual_break = break_mins
                rec.break_minutes = actual_break  # deduct entire break
        elif exceed_rest_col:
            # Only one exceed-rest value available
            exceed = td_to_minutes(ws.cell(row=row_idx, column=exceed_rest_col).value)
            if exceed is not None and exceed > 0:
                rec.break_minutes = exceed
                actual_break = exceed + standard_rest

        # Compute deductible break for structured rest (standard/exceed columns)
        if actual_break > 0 and not simple_break_col:
            rec.break_minutes = engine.compute_deductible_break(
                actual_break, standard_rest, cfg_rest['deduction_mode']
            )

        # Net = gross - deductible break
        if rec.gross_minutes > 0:
            rec.net_minutes = max(0.0, rec.gross_minutes - rec.break_minutes)

        records.append(rec)

    return records


def detect_employment_type(ws) -> Tuple[bool, Optional[float], Optional[float]]:
    """Return (is_full_time, standard_hours_per_day, standard_days_per_month)."""
    has_basic = find_cell_by_keyword(ws, 'BASIC HOURS') is not None
    has_ot = find_cell_by_keyword(ws, 'OT') is not None
    is_full_time = has_basic or has_ot

    std_hours = None
    std_days = None
    basic_cell = find_cell_by_keyword(ws, 'BASIC HOURS')
    if basic_cell:
        parsed = parse_basic_hours_text(str(basic_cell[2]))
        if parsed:
            std_days, total_hours = parsed
            if std_days and total_hours:
                std_hours = total_hours / std_days
                std_days = std_days
    return is_full_time, std_hours, std_days


def process_workbook(input_path: str, output_path: str, year: int = 2026, month: int = 6):
    engine = PayrollEngine()
    wb_in = openpyxl.load_workbook(input_path, data_only=True)

    all_reports = []
    for sheet_name in wb_in.sheetnames:
        ws = wb_in[sheet_name]
        is_full_time, std_hours, std_days = detect_employment_type(ws)
        records = extract_records(ws, year, month, engine)
        if not records:
            continue
        payroll = engine.compute_employee_payroll(
            employee_name=sheet_name,
            is_full_time=is_full_time,
            records=records,
            year=year,
            month=month,
            standard_hours_per_day=std_hours,
            standard_days_per_month=std_days,
        )
        report = engine.format_payroll_report(payroll)
        all_reports.append(report)

    # Write output
    wb_out = openpyxl.Workbook()
    ws_sum = wb_out.active
    ws_sum.title = '薪资汇总'

    headers = [
        '员工', '全职/兼职', '工作天数', '迟到天数', '早退天数', '奇数打卡天数',
        '全勤奖资格', '在司总工时', '净总工时', '扣除休息', '基本工时',
        '加班 OT', '假期工时', '假期标准', '假期加班', '最终加班 OOT',
        '最终工时', '兼职工资（示例）'
    ]
    ws_sum.append(headers)
    for cell in ws_sum[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='DDDDDD')

    for r in all_reports:
        ws_sum.append([
            r['employee_name'],
            '全职' if r['is_full_time'] else '兼职',
            r['work_days'],
            r['late_days'],
            r['early_leave_days'],
            r['odd_punch_days'],
            '是' if r['full_attendance_eligible'] else '否',
            round(r['total_gross_hours'], 2),
            round(r['total_net_hours'], 2),
            round(r['total_break_hours'], 2),
            round(r['basic_hours'], 2),
            round(r['ot_hours'], 2),
            round(r['ph_work_hours'], 2),
            round(r['ph_standard_hours'], 2),
            round(r['ph_ot_hours'], 2),
            round(r['oot_hours'], 2),
            round(r['final_work_hours'], 2),
            r['part_time_wage'],
        ])

    # Add daily details sheet
    ws_det = wb_out.create_sheet('每日明细')
    ws_det.append([
        '员工', '日期', '公共假期', '上班', '下班', '在司分钟', '扣除休息',
        '净工时(小时)', '迟到分钟', '早退分钟'
    ])
    for sheet_name in wb_in.sheetnames:
        ws = wb_in[sheet_name]
        is_full_time, std_hours, std_days = detect_employment_type(ws)
        records = extract_records(ws, year, month, engine)
        if not records:
            continue
        for rec in records:
            if rec.net_minutes <= 0:
                continue
            ws_det.append([
                sheet_name,
                rec.day,
                '是' if rec.is_public_holiday else '否',
                engine.minutes_to_time_str(rec.start_minutes),
                engine.minutes_to_time_str(rec.end_minutes),
                round(rec.gross_minutes, 1),
                round(rec.break_minutes, 1),
                round(rec.net_minutes / 60.0, 2),
                round(rec.late_minutes, 1),
                round(rec.early_minutes, 1),
            ])

    for ws in wb_out.worksheets:
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                try:
                    if cell.value is not None:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

    wb_out.save(output_path)
    print(f'薪资计算结果已保存: {output_path}')


def main(argv=None):
    parser = argparse.ArgumentParser(description='Process JUN.xlsx payroll')
    parser.add_argument('input', help='Input JUN.xlsx')
    parser.add_argument('-o', '--output', default='JUN_薪资计算结果.xlsx', help='Output path')
    parser.add_argument('--year', type=int, default=2026)
    parser.add_argument('--month', type=int, default=6)
    args = parser.parse_args(argv)

    if not Path(args.input).exists():
        print(f'Error: file not found: {args.input}', file=sys.stderr)
        return 1

    process_workbook(args.input, args.output, args.year, args.month)
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
